"""The Low Carbon London release: 5,567 households at half hour resolution, read into DuckDB.

The raw file (167 million rows) is never committed. It is unpacked from the London
Datastore zip, converted once to a parquet file under data/external/london, checked for
duplicated readings and gaps, and every table the meter module needs is aggregated from
that parquet. Household ids never leave data/external: the committed tables are cluster
and group aggregates.
"""

from __future__ import annotations

import shutil
import subprocess
import zipfile
from dataclasses import dataclass
from pathlib import Path

import duckdb
import polars as pl

RAW_ZIP = "LCL-FullData.zip"
RAW_CSV = "CC_LCL-FullData.csv"
READINGS = "readings.parquet"
TARIFFS = "Tariffs.xlsx"


@dataclass(frozen=True)
class LondonChecks:
    rows: int
    households: int
    tou_households: int
    std_households: int
    first: str
    last: str
    null_readings: int
    duplicate_readings: int
    missing_half_hours: int


def connect(
    memory_limit: str = "3GB", threads: int = 2, temp_dir: Path | None = None
) -> duckdb.DuckDBPyConnection:
    """A DuckDB connection with the memory limit below physical memory, so it spills rather than dies."""
    con = duckdb.connect()
    con.execute(f"set memory_limit='{memory_limit}'")
    con.execute(f"set threads={threads}")
    if temp_dir is not None:
        temp_dir.mkdir(parents=True, exist_ok=True)
        con.execute(f"set temp_directory='{temp_dir.as_posix()}'")
    con.execute("set preserve_insertion_order=false")
    return con


def unpack(folder: Path) -> Path:
    """Extract the consolidated CSV from the zip if it is not already on disk.

    The London Datastore zip is compressed with Deflate64, which Python's zipfile refuses, so
    Info-ZIP's unzip does the extraction when it is installed and zipfile is the fallback."""
    csv = folder / RAW_CSV
    if csv.exists():
        return csv
    archive = folder / RAW_ZIP
    if not archive.exists():
        raise FileNotFoundError(
            f"neither {csv} nor {archive} exists; download the release from the London Datastore"
        )
    with zipfile.ZipFile(archive) as zf:
        members = [m for m in zf.namelist() if m.lower().endswith(".csv")]
        if not members:
            raise ValueError(f"{archive} holds no csv member")
        member = members[0]
        method = zf.getinfo(member).compress_type
    if shutil.which("unzip") and method not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
        partial = csv.with_suffix(".csv.part")
        with partial.open("wb") as dst:
            subprocess.run(["unzip", "-p", str(archive), member], stdout=dst, check=True)
        partial.rename(csv)
        return csv
    with zipfile.ZipFile(archive) as zf, zf.open(member) as src, csv.open("wb") as dst:
        while chunk := src.read(1 << 24):
            dst.write(chunk)
    return csv


def to_parquet(folder: Path, con: duckdb.DuckDBPyConnection | None = None) -> Path:
    """The raw CSV as typed parquet: household, tariff, half hour timestamp, kWh. Runs once."""
    out = folder / READINGS
    if out.exists():
        return out
    csv = unpack(folder)
    con = con or connect(temp_dir=folder / "duckdb_tmp")
    # The header's last column carries a trailing space in the release, and missing readings
    # are the word Null, so the columns are read by position and cast with try_cast.
    con.execute(
        f"""
        copy (
            select
                column0 as household,
                case when trim(column1) = 'ToU' then 'ToU' else 'Std' end as tariff,
                try_cast(column2 as timestamp) as ts,
                try_cast(trim(column3) as double) as kwh
            from read_csv('{csv.as_posix()}', header=true, columns={{'column0':'varchar','column1':'varchar','column2':'varchar','column3':'varchar'}}, delim=',', quote='"', ignore_errors=true)
        ) to '{out.as_posix()}' (format parquet, compression zstd, row_group_size 1000000)
        """
    )
    return out


def checks(readings: Path, con: duckdb.DuckDBPyConnection | None = None) -> LondonChecks:
    con = con or connect()
    path = readings.as_posix()
    row = con.execute(
        f"""
        select count(*), count(distinct household),
               count(distinct case when tariff = 'ToU' then household end),
               count(distinct case when tariff = 'Std' then household end),
               min(ts), max(ts), sum(case when kwh is null then 1 else 0 end)
        from read_parquet('{path}')
        """
    ).fetchone()
    assert row is not None
    dup = con.execute(
        f"""
        select coalesce(sum(n - 1), 0) from (
            select household, ts, count(*) as n from read_parquet('{path}') group by 1, 2 having count(*) > 1
        )
        """
    ).fetchone()
    assert dup is not None
    gaps = con.execute(
        f"""
        with spans as (
            select household, min(ts) as first_ts, max(ts) as last_ts, count(distinct ts) as have
            from read_parquet('{path}') group by 1
        )
        select coalesce(sum(cast(date_diff('minute', first_ts, last_ts) / 30 + 1 as bigint) - have), 0) from spans
        """
    ).fetchone()
    assert gaps is not None
    return LondonChecks(
        rows=int(row[0]),
        households=int(row[1]),
        tou_households=int(row[2]),
        std_households=int(row[3]),
        first=str(row[4]),
        last=str(row[5]),
        null_readings=int(row[6]),
        duplicate_readings=int(dup[0]),
        missing_half_hours=int(gaps[0]),
    )


def read_tariffs(folder: Path) -> pl.DataFrame:
    """The 2013 dynamic time of use schedule: one row per half hour with the price band."""
    import openpyxl

    book = openpyxl.load_workbook(folder / TARIFFS, read_only=True, data_only=True)
    sheet = book[book.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    data = [r for r in rows[1:] if r and r[0] is not None]
    frame = pl.DataFrame({header[i]: [r[i] for r in data] for i in range(len(header)) if header[i]})
    time_col = next(c for c in frame.columns if "time" in c.lower() or "date" in c.lower())
    band_col = next(c for c in frame.columns if "tariff" in c.lower() or "price" in c.lower())
    out = frame.select(
        pl.col(time_col).alias("ts"), pl.col(band_col).cast(pl.Utf8).str.strip_chars().alias("band")
    )
    if out["ts"].dtype == pl.Utf8:
        out = out.with_columns(pl.col("ts").str.strptime(pl.Datetime, strict=False))
    return out.with_columns(pl.col("ts").cast(pl.Datetime("us"))).sort("ts")
